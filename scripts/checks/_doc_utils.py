"""
Shared utilities for document processing and manifest loading.

Provides path constants, tracked-change-aware text extraction,
and paragraph building from Word documents.
"""
from pathlib import Path
import openpyxl
from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[2]
MANIFEST_PATH = REPO_ROOT / "output" / "numbers_manifest.xlsx"
DOCX_PATH = REPO_ROOT / "manuscript" / "PRWP working draft - Innovation Patterns in WBG Portfolio.docx"


# ---------------------------------------------------------------------------
# Tracked-change-aware text extraction
# ---------------------------------------------------------------------------

def get_para_text_accepted(para_element) -> str:
    """
    Extract paragraph text as it would appear after accepting all tracked changes:
      - Include text inside <w:ins> elements
      - Exclude text inside <w:del> elements
      - Include plain <w:r> run text
    """
    parts = []
    for node in para_element.iter():
        tag = node.tag.split("}")[-1] if "}" in node.tag else node.tag

        # Skip deleted content
        if tag == "del":
            continue

        # Include text nodes that are direct children of runs (w:r) or insertions (w:ins > w:r)
        if tag == "t":
            # Check ancestors: if any ancestor is a w:del, skip
            in_del = any(
                (p.tag.split("}")[-1] if "}" in p.tag else p.tag) == "del"
                for p in node.iterancestors()
            )
            if not in_del and node.text:
                parts.append(node.text)
    return "".join(parts)


def build_doc_paragraphs(doc: Document) -> list[str]:
    """Return list of accepted-text strings, one per paragraph."""
    return [get_para_text_accepted(p._element) for p in doc.paragraphs]


# ---------------------------------------------------------------------------
# Manifest loading
# ---------------------------------------------------------------------------

def load_manifest() -> list[dict]:
    """Load manifest from numbers_manifest.xlsx and return list of dicts."""
    wb = openpyxl.load_workbook(MANIFEST_PATH)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    cols = rows[0]
    return [dict(zip(cols, row)) for row in rows[1:]]
